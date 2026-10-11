"""레지스터 맵(엑셀 정본, CSV 허용)을 직접 파싱하고 검사한다.

열: register, offset, field, bits, access, reset, description
- 같은 (register, offset)이 레지스터 워드 하나다. 한 레지스터가 여러 워드에 걸칠 수 있다.
- 배열 레지스터는 이름에 `[i]`를 붙이고 offset을 `0x100+0x10*i`처럼 쓴다. 배열 크기는 파라미터에서 가져온다
  (`<이름>_DEPTH` 또는 `array_depth`로 지정).
- 문서(Word/PDF) 안의 레지스터 표는 일부 열이 빠진 사본이므로 입력으로 쓰지 않고, 엑셀과 다르면 경고만 한다.

    python3 -m pipeline.spec.regmap <registers.xlsx|csv> [--spec spec.md] [--array DCT=DAT_DEPTH]
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path

from pipeline.spec.issues import SpecIssue

COLUMNS = ("register", "offset", "field", "bits", "access", "reset", "description")
ACCESS = {"RW", "RO", "WO", "W1C", "W1S", "RC", "RS"}
WORD_BITS = 32


@dataclass(frozen=True)
class Field:
    register: str          # 배열이면 `[i]`를 뗀 이름
    array: bool
    offset: str            # 원문 (식일 수 있음)
    name: str
    hi: int
    lo: int
    access: str
    reset: int | None
    description: str
    row: int               # 원본 행 번호 (머리가 1)

    @property
    def width(self) -> int:
        return self.hi - self.lo + 1


def read_rows(path: Path) -> list[dict[str, str]]:
    """엑셀(.xlsx, 첫 시트) 또는 CSV. 열 이름은 소문자로 맞춘다."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl

        ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = [[("" if v is None else str(v)).strip() for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        with path.open(encoding="utf-8-sig", newline="") as f:
            rows = [[c.strip() for c in r] for r in csv.reader(f)]
    head = [h.lower() for h in rows[0]]
    return [dict(zip(head, r)) for r in rows[1:] if any(r)]


def _int(text: str) -> int | None:
    t = text.strip().replace("_", "")
    try:
        return int(t, 0)
    except ValueError:
        return None


def parse(path: str | Path) -> tuple[list[Field], list[SpecIssue]]:
    path = Path(path)
    rows = read_rows(path)
    issues: list[SpecIssue] = []
    missing = [c for c in COLUMNS if rows and c not in rows[0]]
    if missing:
        return [], [SpecIssue("REG_COLUMNS", "error", path.name, f"필요한 열이 없음: {missing}")]
    out = []
    for n, r in enumerate(rows, start=2):
        where = f"{path.name}:{n}"
        m = re.fullmatch(r"(\d+)(?::(\d+))?", r["bits"])
        if not m:
            issues.append(SpecIssue("REG_BITS", "error", where, f"bits 형식이 틀림: {r['bits']!r} (예: 7:0, 3)"))
            continue
        hi, lo = int(m.group(1)), int(m.group(2) if m.group(2) is not None else m.group(1))
        if hi < lo or hi >= WORD_BITS:
            issues.append(SpecIssue("REG_BITS", "error", where, f"bits 범위가 틀림: {r['bits']} (0~{WORD_BITS - 1}, 높은 쪽이 앞)"))
            continue
        if r["access"].upper() not in ACCESS:
            issues.append(SpecIssue("REG_ACCESS", "error", where, f"알 수 없는 access: {r['access']} (허용: {sorted(ACCESS)})"))
        reset = _int(r["reset"]) if r["reset"] else None
        if r["reset"] and reset is None:
            issues.append(SpecIssue("REG_RESET", "error", where, f"리셋값을 읽을 수 없음: {r['reset']}"))
        elif reset is not None and reset >= (1 << (hi - lo + 1)):
            issues.append(SpecIssue("REG_RESET", "error", where, f"리셋값 {r['reset']}이 필드 폭 {hi - lo + 1}비트를 넘음"))
        name = r["register"]
        array = name.endswith("[i]")
        out.append(Field(name[:-3] if array else name, array, r["offset"], r["field"], hi, lo, r["access"].upper(),
                         reset, r["description"], n))
    return out, issues


def _offsets(f: Field, depth: int | None) -> list[int] | None:
    if not f.array:
        v = _int(f.offset)
        return None if v is None else [v]
    if depth is None:
        depth = 1
    try:
        return [int(eval(f.offset, {"__builtins__": {}}, {"i": i})) for i in range(depth)]  # 식: 0x100+0x10*i
    except Exception:
        return None


def check(fields: list[Field], params: dict[str, int] | None = None, array_depth: dict[str, str] | None = None,
          addr_bits: int | None = None, where: str = "registers") -> list[SpecIssue]:
    """오프셋 정렬, 필드 겹침, 워드 겹침(배열 펼친 뒤), 주소 폭."""
    params, array_depth = params or {}, array_depth or {}
    issues: list[SpecIssue] = []
    words: dict[tuple[str, str], list[Field]] = {}
    for f in fields:
        words.setdefault((f.register, f.offset), []).append(f)
    occupied: dict[int, tuple[str, int]] = {}
    for (reg, off), fs in words.items():
        w = f"{where}:{fs[0].row}"
        for a in range(len(fs)):
            for b in range(a + 1, len(fs)):
                x, y = fs[a], fs[b]
                if x.lo <= y.hi and y.lo <= x.hi:
                    issues.append(SpecIssue("REG_FIELD_OVERLAP", "error", w, f"{reg}의 {x.name}[{x.hi}:{x.lo}]와 {y.name}[{y.hi}:{y.lo}]가 겹침"))
        names = [f.name for f in fs]
        for dup in sorted({n for n in names if names.count(n) > 1}):
            issues.append(SpecIssue("REG_FIELD_DUP", "error", w, f"{reg}에 필드 {dup}가 두 번 있음"))
        depth = None
        if fs[0].array:
            pname = array_depth.get(reg, f"{reg}_DEPTH")
            depth = params.get(pname)
            if depth is None:
                issues.append(SpecIssue("REG_ARRAY_DEPTH", "warning", w, f"배열 {reg}의 크기를 모름 (파라미터 {pname} 없음). 첫 항목만 검사함"))
        offs = _offsets(fs[0], depth)
        if offs is None:
            issues.append(SpecIssue("REG_OFFSET", "error", w, f"offset을 읽을 수 없음: {off}"))
            continue
        for o in offs:
            if o % 4:
                issues.append(SpecIssue("REG_ALIGN", "error", w, f"{reg} 오프셋 0x{o:X}가 4바이트 정렬이 아님"))
            if addr_bits is not None and o >= (1 << addr_bits):
                issues.append(SpecIssue("REG_ADDR_RANGE", "error", w, f"{reg} 오프셋 0x{o:X}가 주소 폭 {addr_bits}비트를 넘음"))
            if o in occupied and occupied[o][0] != reg:
                issues.append(SpecIssue("REG_WORD_OVERLAP", "error", w, f"{reg}와 {occupied[o][0]}가 같은 주소 0x{o:X}를 씀"))
            occupied.setdefault(o, (reg, fs[0].row))
    return issues


# 한글 조사가 바로 붙으므로("CTRL.enable이") \b 대신 ASCII 식별자 경계를 쓴다.
_REF = re.compile(r"(?<![A-Za-z0-9_])([A-Z][A-Z0-9_]*[A-Z0-9])\.([a-z][a-z0-9_]*)(?![A-Za-z0-9_])")
_REGWORD = re.compile(r"(?<![A-Za-z0-9_])([A-Z][A-Z0-9_]*[A-Z0-9])(?:\[[^\]]*\])?\s*레지스터")


def cross_check(spec, fields: list[Field]) -> list[SpecIssue]:
    """SPEC 본문이 인용한 `레지스터.필드`와 `<이름> 레지스터`가 맵에 있는지. 문서 안 레지스터 표와도 대조한다."""
    have = {(f.register, f.name) for f in fields}
    regs = {f.register for f in fields}
    issues: list[SpecIssue] = []
    fname = spec.path.name if spec.path else "spec"
    for sec in spec.sections.values():
        for no, ln in sec.lines:
            for reg, fld in _REF.findall(ln):
                if reg in regs and (reg, fld) not in have:
                    issues.append(SpecIssue("REG_REF_FIELD", "error", f"{fname}:{no}", f"레지스터 맵에 없는 필드: {reg}.{fld}"))
                elif reg not in regs:
                    issues.append(SpecIssue("REG_REF_REG", "warning", f"{fname}:{no}", f"레지스터 맵에 없는 레지스터: {reg} ({reg}.{fld})"))
            for reg in _REGWORD.findall(ln):
                if reg not in regs:  # "<이름> 레지스터"는 추정이라(예: "APB 레지스터로 다룸") 경고로만 낸다
                    issues.append(SpecIssue("REG_REF_REG", "warning", f"{fname}:{no}", f"레지스터 맵에 없는 레지스터일 수 있음: {reg}"))
        for t in sec.tables:
            cols = {h.lower(): h for h in t.header}
            rk = next((cols[k] for k in ("register", "레지스터") if k in cols), None)
            fk = next((cols[k] for k in ("field", "필드") if k in cols), None)
            bk = next((cols[k] for k in ("bits", "비트") if k in cols), None)
            if not (rk and fk):
                continue
            for r, no in zip(t.dicts(), t.row_lines):
                reg = r[rk].removesuffix("[i]")
                f = next((x for x in fields if x.register == reg and x.name == r[fk]), None)
                if f is None:
                    issues.append(SpecIssue("REG_DOC_TABLE", "warning", f"{fname}:{no}", f"문서 표의 {r[rk]}.{r[fk]}가 레지스터 맵에 없음"))
                elif bk and r[bk] and r[bk] != (f"{f.hi}:{f.lo}" if f.hi != f.lo else str(f.hi)):
                    issues.append(SpecIssue("REG_DOC_TABLE", "warning", f"{fname}:{no}",
                                            f"문서 표의 {r[rk]}.{r[fk]} 비트({r[bk]})가 레지스터 맵({f.hi}:{f.lo})과 다름. 레지스터 맵을 따름"))
    return issues


def main(argv=None) -> int:
    from pipeline.spec.l1 import block_top_l1
    from pipeline.spec.parse import load

    ap = argparse.ArgumentParser(description="레지스터 맵 파싱과 검사")
    ap.add_argument("regmap")
    ap.add_argument("--spec")
    ap.add_argument("--array", action="append", default=[], help="배열 크기 파라미터 (예: DCT=DAT_DEPTH)")
    args = ap.parse_args(argv)
    fields, issues = parse(args.regmap)
    params, addr_bits, spec = {}, None, None
    if args.spec:
        spec = load(args.spec)
        ir, _ = block_top_l1(spec)
        params = {p["name"]: int(p["default"]) for p in ir["params"] if str(p["default"]).isdigit()}
        paddr = next((s for p in ir["ports"] if p.get("protocol") == "apb" for s in p["signals"] if s["name"].endswith("addr")), None)
        m = re.search(r"\[(\d+):0\]", paddr["type"]) if paddr else None
        addr_bits = int(m.group(1)) + 1 if m else None
    issues += check(fields, params, dict(a.split("=", 1) for a in args.array), addr_bits, Path(args.regmap).name)
    if spec:
        issues += cross_check(spec, fields)
    print(json.dumps({"fields": len(fields), "registers": len({f.register for f in fields}),
                      "issues": [i.to_dict() for i in issues]}, ensure_ascii=False, indent=1))
    return 1 if any(i.severity == "error" for i in issues) else 0


if __name__ == "__main__":
    sys.exit(main())
